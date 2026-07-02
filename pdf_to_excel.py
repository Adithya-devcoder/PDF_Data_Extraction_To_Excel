"""
PDF Cost Sheet -> Excel Template extractor v4 (append + strict header/notes)

Install once:
    py -m pip install pdfplumber openpyxl

APPEND MODE - reads only NEW PDFs that are not already processed:
    py pdf_to_excel.py --folder "PDFs" --template "Excel Template.xlsx" --out "output.xlsx"

FRESH MODE - rebuilds from the template and reads ALL PDFs from the beginning:
    py pdf_to_excel.py --folder "PDFs" --template "Excel Template.xlsx" --out "output.xlsx" --reset

Important:
- Close output.xlsx before running. Excel locks open files on Windows.
- The hidden Processed_Files sheet prevents duplicate appending.
- Commodity Working, Import Cost Working, and Total Cost are rebuilt from the full Product Cost Sheet every run.
"""
from __future__ import annotations

import argparse
from copy import copy
import re
from pathlib import Path
from typing import Any, Dict, List, Tuple

import pdfplumber
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


# ---------------------------- basic helpers ----------------------------

def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.replace("\u0001", " ").replace("\ufffe", " ")
        value = re.sub(r"\s+", " ", value.replace("\n", " ")).strip()
        return value if value else None
    return value


def to_number(value: Any) -> Any:
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    s = str(value).replace(",", "").strip()
    if s in {"-", "--", "---", ""}:
        return s
    try:
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    except Exception:
        pass
    return value


def safe_float(value: Any) -> float:
    value = to_number(value)
    try:
        return float(value)
    except Exception:
        return 0.0


def extract_text(pdf_path: Path) -> str:
    chunks: List[str] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            chunks.append(page.extract_text(x_tolerance=1, y_tolerance=3) or "")
    return "\n".join(chunks)


def find_between(text: str, left: str, right: str | None = None) -> str | None:
    pattern = re.escape(left) + r"\s*(.*?)" + (re.escape(right) if right else r"(?:\n|$)")
    m = re.search(pattern, text, flags=re.I | re.S)
    return clean(m.group(1)) if m else None


def first_match(text: str, patterns: List[str]) -> str | None:
    for pat in patterns:
        m = re.search(pat, text, flags=re.I | re.S)
        if m:
            return clean(m.group(1))
    return None


# ---------------------------- PDF extraction ----------------------------

def strip_note_side_text(value: str) -> str:
    value = re.sub(
        r"\s*Total\s+(?:LH\s+|RH\s+|Tool\s+)?Cost\s*:?\s*(?:in\s+Rs\.?|Rs\.?|₹)?\s*[\d,.\s-]*",
        " ",
        value,
        flags=re.I,
    )
    return clean(value) or ""


def normalize_note_item(number: int, body: str) -> str:
    if number == 2:
        body = re.sub(r"\s+Volume\s*$", "", body, flags=re.I)
    if number == 4 and re.search(r"\band\s*$", body, flags=re.I):
        body = f"{body} Volume"
    return clean(body) or ""


def extract_notes(full_text: str) -> str | None:
    """Extract only the numbered Notes section, excluding nearby tables/totals."""
    matches = list(re.finditer(r"(?:^|\n)\s*Notes?\s*!?\s*:?\s*", full_text, flags=re.I))
    if not matches:
        return None
    notes = full_text[matches[-1].end():]

    # Stop before unrelated summary/working blocks if they appear after notes.
    stop = re.search(
        r"\n\s*(?:PRODUCT COST SHEET|Project Details|Sl\.?\s*No|Import Cost Working|"
        r"Commodity Working|Final Summary|Prepared|Pre\s+M|Verified|Approved|Authorised|"
        r"Commodity Group|Product Group|Sameas\s*PN|Supplier\s+Curr|C-\s*Casting)\b",
        notes,
        flags=re.I,
    )
    if stop:
        notes = notes[:stop.start()]

    first_number = re.search(r"(?<!\d)1\s*[\).]\s*", notes)
    if first_number:
        notes = notes[first_number.start():]

    markers = list(re.finditer(r"(?<!\d)(\d{1,2})\s*[\).]\s*", notes))
    items: Dict[int, str] = {}
    for i, marker in enumerate(markers):
        number = int(marker.group(1))
        if number < 1 or number > 20:
            continue
        end = markers[i + 1].start() if i + 1 < len(markers) else len(notes)
        body = normalize_note_item(number, strip_note_side_text(notes[marker.end():end]))
        if body:
            items.setdefault(number, body)

    if items:
        notes = " ".join(f"{number}) {items[number]}" for number in sorted(items))
    else:
        notes = strip_note_side_text(notes)

    if not notes:
        return None
    # Keep Excel readable.
    return notes[:3000]


def extract_forex_from_import_working(full_text: str) -> Dict[str, Any]:
    """Fallback exchange rates from import-cost working lines.

    Some newer PDFs omit RMB from the header but include lines like
    "32489814 RMB 11.95 ..." in the import-cost working section.
    """
    rates: Dict[str, Any] = {}
    for line in full_text.splitlines():
        m = re.match(r"\s*\S+\s+(USD|EUR|EURO|GBP|JPY|RMB)\s+(\d+(?:\.\d+)?)\b", line, flags=re.I)
        if not m:
            continue
        curr = normalize_currency(m.group(1))
        if not curr:
            continue
        header_key = "EURO" if curr == "EUR" else curr
        rates.setdefault(header_key, to_number(m.group(2)))
    return rates


def dash_if_blank(value: Any) -> Any:
    value = clean(value)
    return value if value is not None else "-"


def normalize_header_key(value: Any) -> str:
    key = re.sub(r"\s+", " ", str(clean(value) or "")).strip()
    aliases = {
        "filename": "File Name",
        "file name": "File Name",
        "costsheetissueno": "CostSheetIssueNo",
        "cost sheet issue no": "CostSheetIssueNo",
        "cost sheet issue no.": "CostSheetIssueNo",
    }
    return aliases.get(key.lower(), key)


def trim_header_value(value: Any) -> Any:
    value = clean(value)
    if not value:
        return None
    # Stop common leaked table/header text caused by PDF text ordering.
    value = re.split(r"\s+(?:EURO|USD|GBP|JPY|RMB|Drawing\s+reference|Design\s+change|PRODUCT\s+COST\s+SHEET|Sl\.?\s*No)\b", str(value), flags=re.I)[0]
    return clean(value)


def marker_if_present(text: str, pattern: str) -> str:
    """For label-only header fields, do not repeat the label as value."""
    return "✓" if re.search(pattern, text, re.I) else "-"


def norm_cell(value: Any) -> str:
    return str(clean(value) or "").replace("|", " ").strip()


def first_non_empty_between(row: List[Any], start_col: int, end_col: int) -> Any:
    """Return first non-empty cell in row[start_col:end_col]."""
    for c in range(start_col, min(len(row), end_col)):
        v = clean(row[c])
        if v is not None:
            return v
    return None


def extract_first_table(pdf_path: Path) -> List[List[Any]]:
    with pdfplumber.open(str(pdf_path)) as pdf:
        if not pdf.pages:
            return []
        tables = pdf.pages[0].extract_tables({
            "vertical_strategy": "lines",
            "horizontal_strategy": "lines",
            "intersection_tolerance": 5,
            "snap_tolerance": 3,
            "join_tolerance": 3,
        }) or pdf.pages[0].extract_tables() or []
        return tables[0] if tables else []


def group_words_by_line(words: List[Dict[str, Any]], tolerance: float = 6.0) -> List[List[Dict[str, Any]]]:
    lines: List[List[Dict[str, Any]]] = []
    for word in sorted(words, key=lambda w: (float(w["top"]), float(w["x0"]))):
        if not lines:
            lines.append([word])
            continue
        line_top = sum(float(w["top"]) for w in lines[-1]) / len(lines[-1])
        if abs(float(word["top"]) - line_top) <= tolerance:
            lines[-1].append(word)
        else:
            lines.append([word])
    return [sorted(line, key=lambda w: float(w["x0"])) for line in lines]


def words_in_range(line: List[Dict[str, Any]], min_x: float, max_x: float) -> str | None:
    text = " ".join(str(w["text"]) for w in line if min_x <= float(w["x0"]) < max_x)
    return clean(text)


FOREIGN_CURRENCIES = {"EUR", "EURO", "USD", "GBP", "JPY", "RMB"}
ALL_CURRENCIES = FOREIGN_CURRENCIES | {"INR"}


def normalize_currency(value: Any) -> str | None:
    value = clean(value)
    if value is None:
        return None
    curr = str(value).upper().replace("€", "EUR").replace("£", "GBP").replace("$", "USD")
    curr = re.sub(r"[^A-Z]", "", curr)
    if curr == "EURO":
        return "EUR"
    return curr if curr in ALL_CURRENCIES else None


def is_design_marker(value: Any) -> bool:
    value = clean(value)
    return str(value).upper() in {"E", "N"} if value is not None else False


def is_inhouse_marker(value: Any) -> bool:
    value = clean(value)
    if value is None:
        return False
    compact = re.sub(r"[\s_-]+", "", str(value).upper())
    return compact in {"IH", "INHOUSE", "INHSE"}


def normalize_commodity_group(value: Any) -> str | None:
    value = clean(value)
    if value is None:
        return None
    commodity = re.sub(r"\s+", " ", str(value)).strip()
    compact = re.sub(r"[\s_-]+", "", commodity.upper())
    if compact in {"", "0", "IH", "INHOUSE"} or commodity in {"-", "--", "---"}:
        return None
    aliases = {
        "MAJ": "Maj Pressings",
        "MIN": "Min Pressings",
        "IMP": "Imported",
        "IMPORT": "Imported",
        "IMPORTS": "Imported",
        "HARDWARE": "Hardwares",
    }
    return aliases.get(compact, commodity)


def product_row_has_extra_category(raw: List[Any]) -> bool:
    """Detect newer PDFs with Category before Commodity.

    Older PDFs map raw[12] directly to Commodity and raw[13] to E/N.
    Newer PDFs insert Category at raw[12], shifting Commodity to raw[13].
    """
    if len(raw) < 16:
        return False
    category = clean(raw[12])
    shifted_commodity = clean(raw[13])
    shifted_design = clean(raw[14])
    shifted_resp = clean(raw[15])
    if not category or not shifted_commodity:
        return False
    if is_design_marker(shifted_commodity):
        return False
    return is_design_marker(shifted_design) or (
        shifted_design is None and is_inhouse_marker(shifted_resp)
    )


def normalize_product_pdf_row(raw: List[Any]) -> List[Any]:
    raw = list(raw)
    if product_row_has_extra_category(raw):
        del raw[12]
    if len(raw) < 31:
        raw = raw + [None] * (31 - len(raw))
    return [to_number(x) for x in raw[:31]]


def first_number(value: Any) -> Any:
    value = clean(value)
    if value is None:
        return None
    m = re.search(r"-?\d[\d,]*(?:\.\d+)?", str(value))
    return to_number(m.group(0)) if m else None


def split_leading_number(value: Any) -> tuple[Any, Any]:
    value = clean(value)
    if value is None:
        return None, None
    m = re.match(r"\s*(-?\d[\d,]*(?:\.\d+)?)(.*)$", str(value))
    if not m:
        return None, value
    number = to_number(m.group(1))
    rest = clean(m.group(2))
    return number, rest


def numbers_from_words(line: List[Dict[str, Any]], min_x: float, max_x: float) -> List[Any]:
    values = []
    for word in line:
        x0 = float(word["x0"])
        if min_x <= x0 < max_x:
            number = first_number(word["text"])
            if number is not None:
                values.append(number)
    return values


def product_line_from_words(line: List[Dict[str, Any]], cs_no: str | None) -> List[Any] | None:
    tokens = [str(w["text"]) for w in line]
    if len(tokens) < 2:
        return None
    if not re.fullmatch(r"\d{1,3}", tokens[0]):
        return None
    if not re.fullmatch(r"\d+(?:\.\d+)*", tokens[1]):
        return None

    raw: List[Any] = [None] * 31
    raw[0] = to_number(tokens[0])
    raw[1] = to_number(tokens[1])

    same_as = words_in_range(line, 80, 120)
    similar_to = words_in_range(line, 120, 155)
    if same_as is None and similar_to and re.search(r"[A-Z0-9]", str(similar_to), flags=re.I):
        same_as, similar_to = similar_to, None
    raw[2] = same_as
    raw[3] = similar_to
    raw[4] = words_in_range(line, 155, 248)
    raw[5] = first_number(words_in_range(line, 248, 273))
    raw[6] = None
    raw[7] = words_in_range(line, 273, 302)
    raw[8] = words_in_range(line, 302, 326)
    raw[9] = words_in_range(line, 326, 390)
    raw[10] = words_in_range(line, 390, 484)
    raw[11] = words_in_range(line, 484, 499)
    raw[12] = normalize_commodity_group(words_in_range(line, 499, 535))
    raw[13] = words_in_range(line, 535, 548)
    raw[14] = words_in_range(line, 548, 568)
    raw[15] = words_in_range(line, 568, 644)

    raw_material = words_in_range(line, 644, 704)
    raw[16], raw[17] = split_leading_number(raw_material)
    rm_tail = words_in_range(line, 680, 704)
    if raw[17] is None and rm_tail:
        raw[17] = rm_tail

    raw[18] = first_number(words_in_range(line, 704, 736))
    raw[19] = first_number(words_in_range(line, 736, 762))
    raw[20] = first_number(words_in_range(line, 762, 793))

    currency_word = None
    for word in line:
        curr = normalize_currency(word["text"])
        if curr:
            currency_word = word
            raw[22] = curr
            break

    if currency_word:
        curr_x = float(currency_word["x0"])
        before_currency = numbers_from_words(line, 760, curr_x)
        after_currency = numbers_from_words(line, curr_x + 1, 970)
        if before_currency:
            raw[21] = before_currency[-1]
        if raw[22] == "INR":
            if after_currency:
                raw[23] = after_currency[0]
            if len(after_currency) > 1:
                raw[24] = after_currency[1]
            if len(after_currency) >= 2:
                raw[26] = after_currency[-2]
                raw[27] = after_currency[-1]
        else:
            if len(after_currency) >= 3:
                raw[24] = after_currency[0]
                raw[26] = after_currency[-2]
                raw[27] = after_currency[-1]
            elif len(after_currency) >= 2:
                raw[26] = after_currency[-2]
                raw[27] = after_currency[-1]

    raw[29] = first_number(words_in_range(line, 1000, 1026))
    raw[30] = words_in_range(line, 1026, 1165)
    return [cs_no] + normalize_product_pdf_row(raw)


def extract_product_rows_from_text_words(pdf_path: Path, cs_no: str | None) -> List[List[Any]]:
    rows: List[List[Any]] = []
    seen_keys = set()
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            words = page.extract_words(x_tolerance=1, y_tolerance=3)
            for line in group_words_by_line(words, tolerance=3):
                row = product_line_from_words(line, cs_no)
                if not row:
                    continue
                key = (str(row[1]), str(row[2]), str(row[3]), str(row[5]))
                if key in seen_keys:
                    continue
                seen_keys.add(key)
                rows.append(row)
    return rows


def extract_compact_header_from_words(pdf_path: Path) -> Dict[str, Any]:
    """Fallback for compact PDFs where the header has values but no extractable grid."""
    h: Dict[str, Any] = {}
    with pdfplumber.open(str(pdf_path)) as pdf:
        if not pdf.pages:
            return h
        words = [
            w for w in pdf.pages[0].extract_words(x_tolerance=1, y_tolerance=3)
            if float(w.get("top", 0)) < 130
        ]

    rows = []
    for line in group_words_by_line(words):
        left = words_in_range(line, 40, 350)
        middle = words_in_range(line, 350, 540)
        issue_marker = words_in_range(line, 540, 620)
        cost = words_in_range(line, 620, 720)
        rate = words_in_range(line, 720, 840)
        if any([left, middle, issue_marker, cost, rate]):
            rows.append({
                "left": left,
                "middle": middle,
                "issue_marker": issue_marker,
                "cost": cost,
                "rate": rate,
            })

    if len(rows) < 4:
        return h

    h["Customer"] = rows[0]["left"]
    h["CSNo"] = rows[0]["cost"]
    h["USD"] = rows[0]["rate"]

    h["Model"] = rows[1]["left"]
    h["Classification/Spec"] = rows[1]["middle"]
    h["CostSheetIssueNo"] = rows[1]["cost"]
    h["GBP"] = rows[1]["rate"]

    h["SOP"] = rows[2]["left"]
    h["Assy Part Desc"] = rows[2]["middle"]
    h["Date"] = rows[2]["cost"]
    h["JPY"] = rows[2]["rate"]

    h["Volume Per"] = rows[3]["left"]
    h["Assy Part Number"] = rows[3]["middle"]
    h["MaterialPric"] = rows[3]["cost"]
    h["EURO"] = rows[3]["rate"]

    return {k: v for k, v in h.items() if clean(v) is not None}


def find_label_position(table: List[List[Any]], label_regex: str, max_row: int = 8) -> tuple[int, int] | None:
    pat = re.compile(label_regex, re.I)
    for r in range(min(len(table), max_row)):
        for c, value in enumerate(table[r]):
            if pat.search(norm_cell(value)):
                return r, c
    return None


def value_after_label_until(table: List[List[Any]], label_regex: str, stop_col: int | None = None) -> Any:
    pos = find_label_position(table, label_regex)
    if not pos:
        return None
    r, c = pos
    row = table[r]
    end = stop_col if stop_col is not None else min(len(row), c + 5)
    return first_non_empty_between(row, c + 1, end)


def extract_header_from_table(table: List[List[Any]]) -> Dict[str, Any]:
    """Extract header from the real PDF grid.
    This avoids random text-order extraction and prevents appending wrong header values.
    """
    h: Dict[str, Any] = {}
    if not table:
        return h

    def cell(r: int, c: int) -> Any:
        try:
            return clean(table[r][c])
        except Exception:
            return None

    # Left project block: stable in all uploaded PDFs.
    h["Customer"] = cell(2, 2)
    h["Model"] = cell(3, 2)
    h["SOP"] = cell(4, 2)
    h["Volume Per"] = cell(5, 2)

    # These values live between column 4 label and the next block at column 7.
    h["NewModel/Replacem\nent"] = value_after_label_until(table, r"New\s*Model\s*/?\s*Replacement|New\s*Model/Replace", stop_col=7)
    h["Bussiness Periods in\nYr"] = value_after_label_until(table, r"Business|Bussiness.*Period", stop_col=7)
    h["BIL SOB"] = value_after_label_until(table, r"BIL\s*SOB|BILL\s*SOB", stop_col=7)

    # Middle product block: value can be col 9 or 10 depending on old/new PDF.
    h["Product Family"] = value_after_label_until(table, r"Product\s*Family|Product\s*family", stop_col=14)
    h["Classification/Spec"] = value_after_label_until(table, r"Classification\s*/?\s*Spec", stop_col=14)
    h["Assy Part Desc"] = value_after_label_until(table, r"Assy\s*Part\s*Desc", stop_col=14)
    h["Assy Part Number"] = value_after_label_until(table, r"Assy\s*Part\s*Number", stop_col=13)

    # Cost sheet block.
    h["CSNo"] = value_after_label_until(table, r"Cost\s*Sheet\s*Number", stop_col=18)
    h["CostSheetIssueNo"] = value_after_label_until(table, r"Cost\s*Sheet\s*issue|CostSheet issue", stop_col=18)
    h["Date"] = value_after_label_until(table, r"^Date$", stop_col=18)
    h["MaterialPric"] = value_after_label_until(table, r"Material\s*Price\s*Basis", stop_col=18)

    # Exchange rates: label/value pairs.
    fx = {"USD": None, "GBP": None, "JPY": None, "RMB": None, "EURO": None}
    for r in range(min(len(table), 8)):
        row = table[r]
        for c, value in enumerate(row):
            label = norm_cell(value).upper().replace("€", "").replace("$", "").strip()
            key = None
            if label.startswith("USD"):
                key = "USD"
            elif label.startswith("GBP"):
                key = "GBP"
            elif label.startswith("JPY"):
                key = "JPY"
            elif label.startswith("RMB"):
                key = "RMB"
            elif label.startswith("EURO") or label.startswith("EUR"):
                key = "EURO"
            if key:
                for cc in range(c + 1, min(len(row), c + 3)):
                    v = norm_cell(row[cc])
                    if re.fullmatch(r"[0-9]+(?:\.[0-9]+)?", v):
                        fx[key] = v
                        break
    h.update(fx)

    # Label-only/check-box areas. User requested dash when no actual value exists.
    for key in [
        "EDD Scrap", "AL Costing", "AL Bar", "Steel Bar",
        "RM Cost Update", "Design Change Part addition/deletion",
        "Source Price Change", "Proto", "Prod",
    ]:
        h[key] = "-"

    return h


def extract_header(pdf_path: Path) -> Dict[str, Any]:
    """Extract one header row per PDF using header table first, with safe fallbacks."""
    text = extract_text(pdf_path)
    table = extract_first_table(pdf_path)
    header = extract_header_from_table(table)
    compact_header = extract_compact_header_from_words(pdf_path)

    # Safe fallbacks only for truly missing fields.
    first_page = text[:12000]
    fallback = {
        "CSNo": first_match(first_page, [r"\b(CS\s*\d+)\b", r"Cost\s*Sheet\s*Number\s*([A-Z]*\s*\d+)"]),
        "Customer": first_match(first_page, [r"Customer\s+(.+?)\s+Product\s+Family", r"Customer\s+(.+?)\s+Product\s+family"]),
        "Model": first_match(first_page, [r"Model\s+(.+?)\s+(?:New\s+Model|Replacement|Classification)"]),
        "Notes": extract_notes(text),
    }
    for k, v in compact_header.items():
        if not header.get(k):
            header[k] = v
    for k, v in fallback.items():
        if not header.get(k):
            header[k] = v
    for k, v in extract_forex_from_import_working(text).items():
        if not header.get(k) or clean(header.get(k)) == "-":
            header[k] = v
    header["Notes"] = extract_notes(text)
    header["File Name"] = pdf_path.name

    text_fields = {"CostSheetIssueNo", "File Name"}
    out = {}
    for k, v in header.items():
        header_key = normalize_header_key(k)
        if header_key == "File Name":
            out[k] = str(v) if v is not None else "-"
            continue
        value = dash_if_blank(v)
        if header_key in text_fields:
            out[k] = clean(value)
        else:
            out[k] = to_number(value)
    return out


def extract_product_rows(pdf_path: Path, cs_no: str | None) -> List[List[Any]]:
    rows: List[List[Any]] = []
    seen_keys = set()
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            # These are table PDFs, not scanned PDFs. extract_tables is the most stable path.
            tables = page.extract_tables({
                "vertical_strategy": "lines",
                "horizontal_strategy": "lines",
                "intersection_tolerance": 5,
                "snap_tolerance": 3,
                "join_tolerance": 3,
            }) or page.extract_tables() or []
            for table in tables:
                for r in table:
                    if not r or not clean(r[0]):
                        continue
                    if not re.fullmatch(r"\d+", str(clean(r[0]))):
                        continue
                    if len(r) < 2 or not clean(r[1]) or not re.fullmatch(r"\d+(?:\.\d+)*", str(clean(r[1]))):
                        continue
                    r = normalize_product_pdf_row(r)
                    key = (str(r[0]), str(r[1]), str(r[2]), str(r[5]))
                    if key in seen_keys:
                        continue
                    seen_keys.add(key)
                    rows.append([cs_no] + r)
    if not rows:
        rows = extract_product_rows_from_text_words(pdf_path, cs_no)
    return rows


# ---------------------------- workbook helpers ----------------------------

def clear_data(ws, start_row: int) -> None:
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)


def first_empty_row(ws, start_row: int) -> int:
    row = max(start_row, ws.max_row + 1)
    while row >= start_row:
        has_value = any(ws.cell(row=row, column=col).value is not None for col in range(1, ws.max_column + 1))
        if has_value:
            return row + 1
        if row == start_row:
            return row
        row -= 1
    return start_row


def append_rows_no_clear(ws, rows: List[List[Any]], start_row: int) -> int:
    if not rows:
        return 0
    row_num = first_empty_row(ws, start_row)
    for row in rows:
        for col_num, value in enumerate(row, start=1):
            ws.cell(row=row_num, column=col_num, value=value)
        row_num += 1
    return len(rows)


def find_header_column(ws, header_name: str) -> int | None:
    target = normalize_header_key(header_name)
    for col in range(1, ws.max_column + 1):
        if normalize_header_key(ws.cell(1, col).value) == target:
            return col
    return None


def copy_column_format(ws, source_col: int, target_col: int) -> None:
    for row in range(1, max(ws.max_row, 1) + 1):
        source = ws.cell(row, source_col)
        target = ws.cell(row, target_col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.font:
            target.font = copy(source.font)
        if source.fill:
            target.fill = copy(source.fill)
        if source.border:
            target.border = copy(source.border)

    source_letter = get_column_letter(source_col)
    target_letter = get_column_letter(target_col)
    if ws.column_dimensions[source_letter].width:
        ws.column_dimensions[target_letter].width = ws.column_dimensions[source_letter].width


def ensure_header_columns(wb) -> None:
    if "Header" not in wb.sheetnames:
        return
    ws = wb["Header"]

    if find_header_column(ws, "File Name") is None:
        ws.insert_cols(1)
        copy_column_format(ws, 2, 1)
        ws.cell(1, 1, "File Name")

    if find_header_column(ws, "RMB") is None:
        jpy_col = find_header_column(ws, "JPY")
        euro_col = find_header_column(ws, "EURO")
        insert_at = (jpy_col + 1) if jpy_col else (euro_col if euro_col else ws.max_column + 1)
        ws.insert_cols(insert_at)
        copy_from = max(1, insert_at - 1)
        copy_column_format(ws, copy_from, insert_at)
        ws.cell(1, insert_at, "RMB")


def get_processed_sheet(wb):
    if "Processed_Files" in wb.sheetnames:
        ws = wb["Processed_Files"]
    else:
        ws = wb.create_sheet("Processed_Files")
        ws.append(["PDF File", "PDF Path", "Product Rows"])
    ws.sheet_state = "hidden"
    return ws


def processed_pdf_names(wb) -> set[str]:
    if "Processed_Files" not in wb.sheetnames:
        return set()
    ws = wb["Processed_Files"]
    return {str(row[0]) for row in ws.iter_rows(min_row=2, values_only=True) if row and row[0]}


def remove_product_category_column(wb) -> None:
    if "Product Cost Sheet" not in wb.sheetnames:
        return
    ws = wb["Product Cost Sheet"]
    for col in range(1, ws.max_column + 1):
        h1 = str(clean(ws.cell(row=1, column=col).value) or "").strip().lower()
        h2 = str(clean(ws.cell(row=2, column=col).value) or "").strip().lower()
        if h1 == "category" or h2 == "category":
            ws.delete_cols(col, 1)
            print("Removed extra Product Cost Sheet column: Category")
            return


def get_last_product_sno(ws, start_row: int = 3, sno_col: int = 2) -> int:
    max_sno = 0
    for row in range(start_row, ws.max_row + 1):
        value = ws.cell(row=row, column=sno_col).value
        try:
            if value is not None and str(value).strip():
                max_sno = max(max_sno, int(float(str(value).strip())))
        except Exception:
            continue
    return max_sno


def renumber_product_rows(product_rows: List[List[Any]], start_from: int) -> None:
    serial = start_from
    for row in product_rows:
        serial += 1
        row[1] = serial


def normalize_body_style(ws, start_row: int, font_name: str = "Calibri", font_size: int = 11) -> None:
    if ws.max_row < start_row:
        return
    body_font = Font(name=font_name, size=font_size, bold=False)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = body_font
            cell.alignment = body_alignment


def autofit_reasonable(ws) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        values = [str(c.value) for c in col_cells if c.value is not None]
        max_len = max((len(v) for v in values), default=10)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 45)


def clear_workbook_data_for_reset(wb) -> None:
    """Guarantee --reset starts from a clean workbook even if the template/output has old rows."""
    for sheet_name, start_row in {
        "Header": 2,
        "Product Cost Sheet": 3,
        "Commodity Working": 2,
        "Import Cost Working": 2,
        "Total Cost": 2,
    }.items():
        if sheet_name in wb.sheetnames:
            clear_data(wb[sheet_name], start_row)
    if "Processed_Files" in wb.sheetnames:
        del wb["Processed_Files"]


# Product row format after Category removal:
# 0 CSNo, 1 SlNo, 2 Level, 3 SameAs, 4 SimilarTo, 5 PartDesc,
# 6 LH, 7 RH, 8 UOM, 9 Source, 10 Drawing, 11 Engg, 12 Routing,
# 13 Commodity, 14 New/Exist, 15 Resp, 16 SourceName, 17 Norms,
# 18 RMGrade, 19 PriceKg, 20 CostPart, 21 SC, 22 Imp, 23 Currency,
# 24 BO, 25 GST, 26 Transport, 27 Landed, 28 MatLH, 29 MatRH, 30 Tool, 31 Remarks

def read_all_product_rows_from_workbook(wb) -> List[List[Any]]:
    ws = wb["Product Cost Sheet"]
    rows = []
    for excel_row in range(3, ws.max_row + 1):
        values = [ws.cell(excel_row, col).value for col in range(1, 33)]
        if values[0] is None and values[1] is None:
            continue
        rows.append(values)
    return rows


def build_commodity_rows(product_rows: List[List[Any]]) -> List[List[Any]]:
    groups: Dict[Tuple[str, str], Dict[str, float]] = {}
    for r in product_rows:
        cs_no = clean(r[0]) or "UNKNOWN"
        commodity = normalize_commodity_group(r[13])
        if not commodity:
            continue
        key = (str(cs_no), str(commodity))
        d = groups.setdefault(key, {"lh_norms": 0, "rh_norms": 0, "nop_lh": 0, "nop_rh": 0, "landed_lh": 0, "landed_rh": 0})
        lh = safe_float(r[6])
        rh = safe_float(r[7])
        landed = safe_float(r[27])
        mat_lh = safe_float(r[28]) or landed * lh
        mat_rh = safe_float(r[29]) or landed * rh
        d["lh_norms"] += lh
        d["rh_norms"] += rh
        d["nop_lh"] += 1 if lh else 0
        d["nop_rh"] += 1 if rh else 0
        d["landed_lh"] += landed if lh else 0
        d["landed_rh"] += landed if rh else 0
        d["cost_lh"] = d.get("cost_lh", 0) + mat_lh
        d["cost_rh"] = d.get("cost_rh", 0) + mat_rh

    out = []
    for (cs_no, commodity), d in groups.items():
        out.append([
            cs_no, commodity,
            d["lh_norms"], d["rh_norms"], d["nop_lh"], d["nop_rh"],
            d["landed_lh"], d["landed_rh"], d.get("cost_lh", 0), d.get("cost_rh", 0)
        ])
    return out


def get_forex_by_cs(wb) -> Dict[str, Dict[str, Any]]:
    ws = wb["Header"]
    headers = [clean(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    result: Dict[str, Dict[str, Any]] = {}
    for row in range(2, ws.max_row + 1):
        cs = clean(ws.cell(row, idx.get("CSNo", 1)).value)
        if not cs:
            continue
        result[str(cs)] = {
            "USD": ws.cell(row, idx.get("USD", 0)).value if idx.get("USD") else None,
            "GBP": ws.cell(row, idx.get("GBP", 0)).value if idx.get("GBP") else None,
            "JPY": ws.cell(row, idx.get("JPY", 0)).value if idx.get("JPY") else None,
            "EUR": ws.cell(row, idx.get("EURO", 0)).value if idx.get("EURO") else None,
            "EURO": ws.cell(row, idx.get("EURO", 0)).value if idx.get("EURO") else None,
            "RMB": ws.cell(row, idx.get("RMB", 0)).value if idx.get("RMB") else None,
        }
    return result


def build_import_rows(product_rows: List[List[Any]], forex_by_cs: Dict[str, Dict[str, Any]]) -> List[List[Any]]:
    out = []
    for r in product_rows:
        cs_no = clean(r[0]) or "UNKNOWN"
        curr_s = normalize_currency(r[23])
        if not curr_s or curr_s == "INR":
            continue
        fx = forex_by_cs.get(str(cs_no), {}).get(curr_s)
        base = r[22] if safe_float(r[22]) else r[24]  # import cost if present, else BO cost
        inr_value = safe_float(base) * safe_float(fx)
        out.append([
            cs_no, r[3], r[4], r[16], curr_s, fx, base, inr_value,
            None, None, None, None, None, None, None, None, None, None, None, None, None, r[27]
        ])
    return out


def build_total_rows(product_rows: List[List[Any]]) -> List[List[Any]]:
    """Build totals only from values available in the extracted product table.

    Do not invent probability values. If the PDF does not explicitly provide
    probability, the probability columns are left blank.
    """
    totals: Dict[str, Dict[str, float]] = {}
    for r in product_rows:
        cs_no = str(clean(r[0]) or "UNKNOWN")
        d = totals.setdefault(cs_no, {"lh": 0, "rh": 0, "tool": 0, "nop_lh": 0, "nop_rh": 0, "landed_lh": 0, "landed_rh": 0})
        lh = safe_float(r[6])
        rh = safe_float(r[7])
        landed = safe_float(r[27])
        mat_lh = safe_float(r[28])
        mat_rh = safe_float(r[29])
        tool = safe_float(r[30])
        d["lh"] += mat_lh
        d["rh"] += mat_rh
        d["landed_lh"] += landed if lh else 0
        d["landed_rh"] += landed if rh else 0
        d["tool"] += tool
        d["nop_lh"] += 1 if lh else 0
        d["nop_rh"] += 1 if rh else 0
    return [[cs, d["lh"], d["rh"], d["tool"], d["nop_lh"], d["landed_lh"], d["lh"], d["nop_rh"], d["landed_rh"], d["rh"], None, None] for cs, d in totals.items()]


def rebuild_derived_sheets(wb) -> Tuple[int, int, int]:
    product_rows = read_all_product_rows_from_workbook(wb)
    forex_by_cs = get_forex_by_cs(wb)

    commodity_rows = build_commodity_rows(product_rows)
    import_rows = build_import_rows(product_rows, forex_by_cs)
    total_rows = build_total_rows(product_rows)

    for sheet_name, rows in [
        ("Commodity Working", commodity_rows),
        ("Import Cost Working", import_rows),
        ("Total Cost", total_rows),
    ]:
        ws = wb[sheet_name]
        clear_data(ws, 2)
        for row in rows:
            ws.append(row)

    return len(commodity_rows), len(import_rows), len(total_rows)


# ---------------------------- main pipeline ----------------------------

def fill_workbook(template: Path, pdf_paths: List[Path], output: Path, reset: bool = False) -> None:
    if output.exists() and not reset:
        wb = load_workbook(output)
        print(f"Using existing workbook: {output}")
    else:
        wb = load_workbook(template)
        print(f"Fresh run from template: {template}")
        clear_workbook_data_for_reset(wb)
        if reset and output.exists():
            print(f"Reset mode: existing output will be overwritten: {output}")

    remove_product_category_column(wb)
    ensure_header_columns(wb)
    processed_ws = get_processed_sheet(wb)
    processed = set() if reset else processed_pdf_names(wb)

    new_headers = []
    new_products = []
    actually_processed = []

    for pdf in pdf_paths:
        pdf = pdf.resolve()
        if pdf.name in processed:
            print(f"Skipping already processed PDF: {pdf.name}")
            continue

        print(f"Processing: {pdf.name}")
        header = extract_header(pdf)
        cs_no = header.get("CSNo")
        products = extract_product_rows(pdf, cs_no)

        new_headers.append(header)
        new_products.extend(products)
        actually_processed.append((pdf.name, str(pdf), len(products)))

    if new_products:
        product_ws = wb["Product Cost Sheet"]
        last_sno = get_last_product_sno(product_ws, start_row=3, sno_col=2)
        renumber_product_rows(new_products, last_sno)

        header_ws = wb["Header"]
        header_cols = [normalize_header_key(c.value) for c in header_ws[1]]
        header_rows = []
        for h in new_headers:
            h_norm = {normalize_header_key(k): v for k, v in h.items()}
            header_rows.append([h_norm.get(col, "-") for col in header_cols])
        append_rows_no_clear(header_ws, header_rows, 2)

        append_rows_no_clear(product_ws, new_products, 3)

        for item in actually_processed:
            processed_ws.append(list(item))
    else:
        print("No new Product Cost rows found to append.")

    # Important fix: rebuild all summary sheets from the complete Product Cost Sheet.
    # This prevents partial/missing Commodity, Import, and Total data after append runs.
    commodity_count, import_count, total_count = rebuild_derived_sheets(wb)

    sheet_start_rows = {
        "Header": 2,
        "Product Cost Sheet": 3,
        "Commodity Working": 2,
        "Import Cost Working": 2,
        "Total Cost": 2,
    }
    for sheet_name, start_row in sheet_start_rows.items():
        ws = wb[sheet_name]
        ws.freeze_panes = "A3" if sheet_name == "Product Cost Sheet" else "A2"
        normalize_body_style(ws, start_row)
        autofit_reasonable(ws)

    wb.save(output)
    print(f"Done. Wrote: {output}")
    print(f"PDFs newly processed: {len(actually_processed)} | New product rows extracted: {len(new_products)}")
    print(f"Rebuilt sheets -> Commodity rows: {commodity_count} | Import rows: {import_count} | Total rows: {total_count}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pdf", type=Path, help="Single PDF path")
    parser.add_argument("--folder", type=Path, help="Folder containing PDFs, example: PDFs")
    parser.add_argument("--template", type=Path, required=True, help="Excel template path")
    parser.add_argument("--out", type=Path, required=True, help="One Excel output file to keep using/appending")
    parser.add_argument("--reset", action="store_true", help="Fresh run: recreate output from template and read all PDFs again")
    args = parser.parse_args()

    if args.folder:
        pdf_paths = sorted(args.folder.glob("*.pdf"))
    elif args.pdf:
        pdf_paths = [args.pdf]
    else:
        raise SystemExit("Pass either --pdf or --folder")

    if not pdf_paths:
        raise SystemExit("No PDFs found")

    fill_workbook(args.template, pdf_paths, args.out, reset=args.reset)


if __name__ == "__main__":
    main()
